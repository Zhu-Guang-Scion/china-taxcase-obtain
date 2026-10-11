#!/usr/bin/env python3
# ============================================================
# Copyright (c) 2026 Zhu-Guang-Scion <Zhu-Guang-Scion@users.noreply.github.com>
# China_Law 法规库与税务案例库工具链 · github.com/Zhu-Guang-Scion · 水印 CLAW-ZGS-2026
# MIT License（见仓库 LICENSE）；数据不入库，仅元数据索引
# ============================================================

"""最终整理与 Markdown 转换管线。



1. 解析清单+进度 → 每 bbbs/url 的现行文件（最新 ok 记录且文件在盘）

2. 复制到 lawlibrary/<库>/<类别>/原文件/

3. 转 md → lawlibrary/<库>/<类别>/Markdown/（front matter+正文）

4. 涉税/涉会计专题 → lawlibrary/专题/涉税法规|涉会计法规/<库>/<类别>/

5. 总清单 lawlibrary/manifest.csv



用法: python organize_md.py [--only 类别] [--no-topic] [--no-md]

"""
__author__ = "Zhu-Guang-Scion"
__copyright__ = "Copyright (c) 2026 Zhu-Guang-Scion (github.com/Zhu-Guang-Scion), MIT"

import csv

import glob

import json

import os

import re

import shutil

import sys



ROOT = os.environ.get("CHINA_LAW_WORK") or os.path.join(os.getcwd(), "work")

LISTS = os.path.join(ROOT, "lists")

FLK_DL = os.path.join(ROOT, "downloads", "flk")

GOV_DL = os.path.join(ROOT, "downloads", "gov")

LIB = os.path.abspath(os.path.join(ROOT, "..", "lawlibrary"))



TAX_KW = ["税收", "税务", "增值税", "消费税", "营业税", "企业所得税", "个人所得税", "关税",

          "车船税", "印花税", "契税", "土地增值税", "房产税", "资源税", "环境保护税",

          "耕地占用税", "烟叶税", "船舶吨税", "车辆购置税", "城市维护建设税", "发票",

          "出口退税", "国际税收", "涉税", "税收征收管理"]

ACC_KW = ["会计", "审计", "注册会计师", "总会计师", "财务"]



FLK_CATS = ["宪法", "法律", "行政法规", "监察法规", "地方法规", "司法解释"]

GOV_CATS = ["部门规章", "地方政府规章"]





def is_tax(title):

    return any(k in title for k in TAX_KW)





def is_acc(title):

    return any(k in title for k in ACC_KW)





def safe_name(name):

    name = re.sub(r'[\\/:*?"<>|\r\n\t]', "_", name).strip().rstrip(".")

    return name[:100] or "unnamed"





TIAO_HEAD = re.compile(r"^第[一二三四五六七八九十百千零\d]+条")

ZHANG_HEAD = re.compile(r"^第[一二三四五六七八九十百千零\d]+[章节编]")





def docx_to_md(path):

    from docx import Document

    doc = Document(path)

    lines = []

    for p in doc.paragraphs:

        t = p.text.strip()

        if not t:

            continue

        style = (p.style.name or "").lower()

        if "heading" in style:

            m = re.search(r"(\d+)", style)

            lvl = int(m.group(1)) if m else 2

            lines.append("#" * min(lvl + 1, 6) + " " + t)

            continue

        if ZHANG_HEAD.match(t) and len(t) <= 30:

            lines.append("## " + t)

            continue

        if TIAO_HEAD.match(t):

            lines.append("**" + t + "**")

            continue

        lines.append(t)

    for tbl in doc.tables:

        lines.append("")

        for i, row in enumerate(tbl.rows):

            cells = [c.text.strip().replace("\n", " ") for c in row.cells]

            lines.append("| " + " | ".join(cells) + " |")

            if i == 0:

                lines.append("|" + "---|" * len(cells))

    return "\n\n".join(lines)





def pdf_to_md(path):

    import pymupdf

    doc = pymupdf.open(path)

    parts = [page.get_text() for page in doc]

    doc.close()

    text = "\n".join(parts)

    out_lines = []

    buf = ""

    for ln in text.splitlines():

        ln = ln.strip()

        if not ln:

            if buf:

                out_lines.append(buf)

                buf = ""

            continue

        buf += ln

        if re.search(r"[。；：:！!？?]$", ln) or TIAO_HEAD.match(ln):

            out_lines.append(buf)

            buf = ""

    if buf:

        out_lines.append(buf)

    return "\n\n".join(

        ("**" + l + "**") if TIAO_HEAD.match(l) else l for l in out_lines if l)





def doc_to_md(path):

    """Word COM 转换 .doc（需本机 MS Word）。"""

    try:

        import win32com.client

        import pythoncom

        pythoncom.CoInitialize()

        word = win32com.client.Dispatch("Word.Application")

        word.Visible = False

        try:

            d = word.Documents.Open(os.path.abspath(path), ReadOnly=True)

            text = d.Content.Text

            d.Close(False)

        finally:

            word.Quit()

            pythoncom.CoUninitialize()

        text = text.replace("\r", "\n")

        lines = [l.strip() for l in text.split("\n") if l.strip()]

        return "\n\n".join(

            ("**" + l + "**") if TIAO_HEAD.match(l) else l for l in lines)

    except Exception as e:

        return f"<!-- .doc 转换失败: {e} -->"





def to_md(src_path):

    ext = os.path.splitext(src_path)[1].lower()

    if ext == ".docx":

        return docx_to_md(src_path)

    if ext == ".pdf":

        return pdf_to_md(src_path)

    if ext == ".doc":

        return doc_to_md(src_path)

    return ""





def norm_title(t):

    return re.sub(r'[《》〈〉\s（）()\-—_",.。·、:：;；!！?？]', '', str(t or ''))





_FALLBACK = None





def fallback_content(title):

    """扫描件PDF/老doc的正文回退：HuggingFace 2023-09 快照（按标题精确匹配）。"""

    global _FALLBACK

    if _FALLBACK is None:

        fp = os.path.join(LISTS, "content_fallback.json")

        _FALLBACK = json.load(open(fp, encoding="utf-8")) if os.path.exists(fp) else {}

    hit = _FALLBACK.get(norm_title(title))

    if hit and hit.get("content"):

        c = str(hit["content"])

        paras = [p.strip() for p in re.split(r"\n+", c) if p.strip()]

        body = "\n\n".join(

            ("**" + p + "**") if TIAO_HEAD.match(p) else p for p in paras)

        note = ("\n\n> 注：原文件为扫描版PDF/旧版Word，正文取自 HuggingFace "

                "twang2218/chinese-law-and-regulations 2023-09 快照（与官方版本标题/日期一致）。")

        return body + note

    return ""





def md_doc(meta, body):

    fm = [f"{k}: {v}" for k, v in

          [("title", meta.get("标题", "")), ("库", meta.get("库", "")),

           ("类别", meta.get("类别", "")), ("公布机关", meta.get("机关", "")),

           ("公布日期", meta.get("公布日期", "")), ("时效性", "现行有效"),

           ("来源", meta.get("来源", "")), ("原文链接", meta.get("链接", ""))]

          if v]

    return "---\n" + "\n".join(fm) + "\n---\n\n" + body + "\n"





def load_latest(path, key):

    latest = {}

    if os.path.exists(path):

        for line in open(path, encoding="utf-8"):

            line = line.strip()

            if line:

                try:

                    rec = json.loads(line)

                    latest[rec[key]] = rec

                except Exception:

                    pass

    return latest





def resolve_flk_current(only):

    result = {}

    for cat in FLK_CATS:

        if only and cat != only:

            continue

        lp = os.path.join(LISTS, f"flk_{cat}.jsonl")

        if not os.path.exists(lp):

            continue

        items = [json.loads(l) for l in open(lp, encoding="utf-8") if l.strip()]

        latest = load_latest(os.path.join(LISTS, f"flk_dl_progress_{cat}.jsonl"), "bbbs")

        d = os.path.join(FLK_DL, cat)

        disk = {os.path.basename(p): p for p in

                glob.glob(os.path.join(d, "*.docx")) + glob.glob(os.path.join(d, "*.pdf"))

                + glob.glob(os.path.join(d, "*.doc"))}

        for it in items:

            rec = latest.get(it["bbbs"])

            fp = disk.get(rec.get("file", "")) if rec and rec.get("ok") else None

            if not fp:

                continue

            result[it["bbbs"]] = {

                "库": "国家法律法规数据库", "类别": cat, "标题": it["title"],

                "机关": it.get("zdjgName", ""), "公布日期": it.get("gbrq", ""),

                "链接": f"https://flk.npc.gov.cn/detail?id={it['bbbs']}",

                "来源": rec.get("source", "官方API"), "src_path": fp,

            }

    return result





def resolve_gov_current(only):

    result = {}

    for cat in GOV_CATS:

        if only and cat != only:

            continue

        lp = os.path.join(LISTS, f"gov_{cat}.jsonl")

        if not os.path.exists(lp):

            continue

        items = [json.loads(l) for l in open(lp, encoding="utf-8") if l.strip()]

        latest = load_latest(os.path.join(LISTS, f"gov_dl_progress_{cat}.jsonl"), "url")

        d = os.path.join(GOV_DL, cat)

        disk = {os.path.basename(p): p for p in glob.glob(os.path.join(d, "*.docx"))}

        for it in items:

            rec = latest.get(it["url"])

            fp = disk.get(rec.get("file", "")) if rec and rec.get("ok") else None

            if not fp:

                continue

            result["gov:" + it["url"]] = {

                "库": "国家规章库", "类别": cat, "标题": it["title"],

                "机关": it.get("org", ""), "公布日期": it.get("date", ""),

                "链接": it["url"], "来源": rec.get("source", "gov.cn详情页"),

                "src_path": fp,

            }

    return result





def main():

    only = None

    if "--only" in sys.argv:

        only = sys.argv[sys.argv.index("--only") + 1]

    do_topic = "--no-topic" not in sys.argv

    do_md = "--no-md" not in sys.argv



    all_items = {}

    all_items.update(resolve_flk_current(only))

    all_items.update(resolve_gov_current(only))

    print(f"现行文件解析: {len(all_items)}", flush=True)



    rows = []

    tax_n = acc_n = md_ok = md_fail = 0

    for key, meta in all_items.items():

        lib, cat = meta["库"], meta["类别"]

        base_name = safe_name(f"{meta['标题']}_{meta['公布日期']}")

        orig_dir = os.path.join(LIB, lib, cat, "原文件")

        os.makedirs(orig_dir, exist_ok=True)

        ext = os.path.splitext(meta["src_path"])[1].lower()

        dst = os.path.join(orig_dir, base_name + ext)

        if not os.path.exists(dst):

            shutil.copy2(meta["src_path"], dst)



        md_path = ""

        if do_md:

            md_dir = os.path.join(LIB, lib, cat, "Markdown")

            os.makedirs(md_dir, exist_ok=True)

            md_path = os.path.join(md_dir, base_name + ".md")

            if not os.path.exists(md_path):

                try:

                    body = to_md(meta["src_path"])

                except Exception:

                    body = ""

                if (not body or len(body.strip()) <= 100) and                         os.path.splitext(meta["src_path"])[1].lower() in ('.docx', '.doc'):

                    body = docx_xml_fallback(meta["src_path"])

                if not body or len(body.strip()) <= 100:

                    body = fallback_content(meta["标题"])

                if body and len(body.strip()) > 100:

                    with open(md_path, "w", encoding="utf-8") as f:

                        f.write(md_doc(meta, body))

                else:

                    md_path = ""

        if md_path:

            md_ok += 1

        else:

            md_fail += 1



        topics = []

        if do_topic:

            if is_tax(meta["标题"]):

                topics.append("涉税")

            if is_acc(meta["标题"]):

                topics.append("涉会计")

            for tp in topics:

                td = os.path.join(LIB, "专题", f"{tp}法规", lib, cat)

                os.makedirs(td, exist_ok=True)

                if not os.path.exists(os.path.join(td, base_name + ext)):

                    shutil.copy2(dst, os.path.join(td, base_name + ext))

                if do_md and md_path:

                    tmd = os.path.join(td, base_name + ".md")

                    if not os.path.exists(tmd):

                        shutil.copy2(md_path, tmd)

        if "涉税" in topics:

            tax_n += 1

        if "涉会计" in topics:

            acc_n += 1



        rows.append({

            "库": lib, "类别": cat, "标题": meta["标题"], "公布机关": meta["机关"],

            "公布日期": meta["公布日期"], "时效性": "现行有效",

            "专题": "、".join(topics), "来源": meta["来源"],

            "原文件": os.path.relpath(dst, LIB),

            "Markdown": os.path.relpath(md_path, LIB) if md_path else "",

            "原文链接": meta["链接"],

        })

        if len(rows) % 2000 == 0:

            print(f"  已处理 {len(rows)}", flush=True)



    os.makedirs(LIB, exist_ok=True)

    csv_path = os.path.join(LIB, "manifest.csv")

    with open(csv_path, "w", encoding="utf-8-sig", newline="") as f:

        w = csv.DictWriter(f, fieldnames=list(rows[0].keys()))

        w.writeheader()

        w.writerows(rows)

    print(f"完成: 条目 {len(rows)}, md成功 {md_ok} 失败 {md_fail}, 涉税 {tax_n}, 涉会计 {acc_n}")

    print(f"-> {csv_path}")





if __name__ == "__main__":

    main()

