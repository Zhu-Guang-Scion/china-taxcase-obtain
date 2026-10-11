#!/usr/bin/env python3
# ============================================================
# Copyright (c) 2026 Zhu-Guang-Scion <Zhu-Guang-Scion@users.noreply.github.com>
# China_Law 法规库与税务案例库工具链 · github.com/Zhu-Guang-Scion · 水印 CLAW-ZGS-2026
# MIT License（见仓库 LICENSE）；数据不入库，仅元数据索引
# ============================================================
__author__ = "Zhu-Guang-Scion"
__copyright__ = "Copyright (c) 2026 Zhu-Guang-Scion (github.com/Zhu-Guang-Scion), MIT"
import os, zipfile, re
from docx import Document

CJK = re.compile(r'[\u4e00-\u9fff]')
TIAO = re.compile(r'第[一二三四五六七八九十百千零\d]+条')

checks = [
    ('lawlibrary/国家规章库/地方政府规章/原文件/厦门市城市生活垃圾管理办法_2004-12-14.docx', 'API重建docx'),
    ('lawlibrary/国家规章库/地方政府规章/Markdown/厦门市城市生活垃圾管理办法_2004-12-14.md', 'API重建md'),
    ('lawlibrary/国家法律法规数据库/司法解释/Markdown/人民检察院检察建议工作规定_2019-02-26.md', 'PUA清洗md'),
    ('lawlibrary/国家法律法规数据库/司法解释/原文件/人民检察院检察建议工作规定_2019-02-26.docx', 'PUA清洗docx'),
]
for fp, kind in checks:
    if not os.path.exists(fp):
        print(kind, '不存在!')
        continue
    if fp.endswith('.docx'):
        z = zipfile.ZipFile(fp)
        ok = 'word/document.xml' in z.namelist()
        doc = Document(fp)
        t = '\n'.join(p.text for p in doc.paragraphs)
    else:
        t = open(fp, encoding='utf-8').read()
        ok = True
    fffd = t.count(chr(0xFFFD))
    pua = len([c for c in t if 0xE000 <= ord(c) <= 0xF8FF])
    gjk = t.count('锟斤拷')
    cjk = len(CJK.findall(t))
    tiao = len(TIAO.findall(t))
    print('%s %s: zipOK=%s 字数=%d FFFD=%d PUA=%d 锟斤拷=%d 条数=%d' % (
        kind, os.path.basename(fp)[:42], ok, cjk, fffd, pua, gjk, tiao))

fp = 'lawlibrary/国家规章库/地方政府规章/原文件/河北省唐山市征收土地地上附着物补偿标准暂行规定_2013-03-14.docx'
if os.path.exists(fp):
    t = '\n'.join(p.text for p in Document(fp).paragraphs)
    print('唐山docx: 锟斤拷=%d 字数=%d' % (t.count('锟斤拷'), len(CJK.findall(t))))
