#!/usr/bin/env python3
# ============================================================
# Copyright (c) 2026 Zhu-Guang-Scion <Zhu-Guang-Scion@users.noreply.github.com>
# China_Law 法规库与税务案例库工具链 · github.com/Zhu-Guang-Scion · 水印 CLAW-ZGS-2026
# MIT License（见仓库 LICENSE）；数据不入库，仅元数据索引
# ============================================================
"""国家规章库(gov.cn) 批量下载器：抓取详情页 HTML，提取正文，转 docx，另存原始 HTML。

用法: python gov_download.py <部门规章|地方政府规章|all> [workers]
断点续下：work/lists/gov_dl_progress_<cat>.jsonl (url -> 状态)
"""
__author__ = "Zhu-Guang-Scion"
__copyright__ = "Copyright (c) 2026 Zhu-Guang-Scion (github.com/Zhu-Guang-Scion), MIT"
import html as htmllib
import json
import os
import random
import re
import sys
import threading
import time

import requests
from docx import Document

sys.path.insert(0, os.path.dirname(__file__))
import govlib

ROOT = os.path.join(os.path.dirname(__file__), "..")
LISTS = os.path.join(ROOT, "lists")
DL_ROOT = os.path.join(ROOT, "downloads", "gov")

_print_lock = threading.Lock()


def log(msg):
    with _print_lock:
        print(msg, flush=True)


def make_session():
    s = requests.Session()
    s.headers.update({
        "User-Agent": "Mozilla/5.0 (Windows NT 10.0; Win64; x64) AppleWebKit/537.36 "
                      "(KHTML, like Gecko) Chrome/131.0.0.0 Safari/537.36",
        "Accept": "text/html,application/xhtml+xml,application/xml;q=0.9,*/*;q=0.8",
        "Accept-Language": "zh-CN,zh;q=0.9",
    })
    return s


def sanitize(name: str) -> str:
    name = re.sub(r'[\\/:*?"<>|\r\n\t]', "_", name)
    return name.strip().rstrip(".")[:100] or "unnamed"


_STYLE_SCRIPT_RE = re.compile(r"<(style|script)\b[^>]*>.*?</\1>", re.S | re.I)


def extract_body(html: str) -> str:
    """提取正文 div（pages_content / trs_editor_view），保留内部 HTML。先剥离 style/script。"""
    html = _STYLE_SCRIPT_RE.sub("", html)
    start = html.find('class="pages_content"')
    if start < 0:
        start = html.find('class="trs_editor_view')
    if start < 0:
        return ""
    pos = html.find(">", start) + 1
    depth = 1
    while depth > 0 and pos < len(html):
        next_open = html.find("<div", pos)
        next_close = html.find("</div>", pos)
        if next_close < 0:
            break
        if 0 <= next_open < next_close:
            depth += 1
            pos = next_open + 4
        else:
            depth -= 1
            if depth == 0:
                return html[html.find(">", start) + 1: next_close]
            pos = next_close + 6
    return ""


_CJK_RE = re.compile(r"[\u4e00-\u9fff]")


def extract_body_generic(html: str) -> str:
    """通用回退：收集全文档中含足量中文的 <p> 块（适用于地方政府官网异构页面）。"""
    parts = []
    for m in _BLOCK_RE.finditer(html):
        tag = m.group("tag").lower()
        if tag not in ("p", "li"):
            continue
        inner = re.sub(r"<[^>]+>", "", m.group("inner"))
        inner = htmllib.unescape(inner).strip()
        cjk = len(_CJK_RE.findall(inner))
        if cjk >= 10 and len(inner) >= 15:
            parts.append(f"<p>{inner}</p>")
    return "\n".join(parts)


_BLOCK_RE = re.compile(
    r"<(?P<tag>p|h[1-6]|li|tr|blockquote)\b[^>]*>(?P<inner>.*?)</(?P=tag)>",
    re.S | re.I)
_TAG_RE = re.compile(r"<[^>]+>")
_WS_RE = re.compile(r"[ \t\xa0]+")


def body_to_blocks(body_html: str):
    """把正文 HTML 拆成 (kind, text) 块序列。"""
    blocks = []
    for m in _BLOCK_RE.finditer(body_html):
        tag = m.group("tag").lower()
        inner = m.group("inner")
        # <br> 转换行
        inner = re.sub(r"<br\s*/?>", "\n", inner, flags=re.I)
        text = _TAG_RE.sub("", inner)
        text = htmllib.unescape(text)
        text = _WS_RE.sub(" ", text)
        lines = [ln.strip() for ln in text.split("\n")]
        lines = [ln for ln in lines if ln]
        if not lines:
            continue
        kind = "h" if tag.startswith("h") else ("li" if tag == "li" else "p")
        for ln in lines:
            blocks.append((kind, ln))
    # 去掉空块
    blocks = [(k, t) for k, t in blocks if t.strip()]
    # 若几乎没解析出块（无 p 标签），退化为纯文本
    if len(blocks) < 3:
        text = _WS_RE.sub(" ", htmllib.unescape(_TAG_RE.sub("\n", body_html)))
        blocks = [("p", ln.strip()) for ln in text.split("\n") if ln.strip()]
    return blocks


def make_docx(title: str, blocks, meta: dict, fpath: str):
    doc = Document()
    doc.add_heading(title, level=1)
    m = [f"{k}：{v}" for k, v in
         [("类别", meta.get("cat", "")), ("公布机关", meta.get("org", "")),
          ("公布日期", meta.get("date", "")), ("文号", meta.get("doc_num", "")),
          ("来源", meta.get("url", ""))] if v]
    p = doc.add_paragraph(" ｜ ".join(m))
    for run in p.runs:
        run.italic = True
        run.font.size = None
    for kind, text in blocks:
        if kind == "h":
            h = doc.add_heading("", level=2)
            r = h.add_run(text)
            r.bold = True
        elif kind == "li":
            doc.add_paragraph(text, style="List Bullet")
        else:
            # 条款加粗
            para = doc.add_paragraph()
            r = para.add_run(text)
            if re.match(r"^第[一二三四五六七八九十百千零\d]+[条章节编]", text):
                r.bold = True
    doc.save(fpath)


class Progress:
    def __init__(self, path):
        self.path = path
        self.lock = threading.Lock()
        self.data = {}
        if os.path.exists(path):
            with open(path, encoding="utf-8") as f:
                for line in f:
                    line = line.strip()
                    if line:
                        try:
                            rec = json.loads(line)
                            self.data[rec["url"]] = rec
                        except Exception:
                            continue

    def mark(self, url, **kv):
        with self.lock:
            rec = {"url": url}
            rec.update(kv)
            self.data[url] = rec
            with open(self.path, "a", encoding="utf-8") as f:
                f.write(json.dumps(rec, ensure_ascii=False) + "\n")

    def done_ok(self, url):
        with self.lock:
            r = self.data.get(url)
            return bool(r and r.get("ok"))


_api_session = None
_api_lock = threading.Lock()


def fetch_content_via_api(item, retries=3):
    """规章库搜索接口按标题取正文字段 f_202321758948（HTML 提取失败的回退）。"""
    global _api_session
    with _api_lock:
        if _api_session is None:
            _api_session = govlib.make_session()
    title = re.sub(r'[《》<>〈〉]', "", item["title"]).strip()
    body = {
        "code": "18258ab0ac9",
        "searchFields": [
            {"fieldName": govlib.F_TITLE, "withHighLight": False, "searchWord": title},
            {"fieldName": govlib.F_TYPE, "searchType": "TERM",
             "withHighLight": False, "searchWord": item["cat"]},
        ],
        "sorts": [{}, {"sortField": govlib.F_DATE, "sortOrder": "DESC"}],
        "resultFields": [govlib.F_TITLE, govlib.F_URL, "f_202321758948"],
        "trackTotalHits": "true", "tableName": "t_1860c735d31",
        "pageSize": 10, "pageNo": 1, "granularity": "ALL",
    }
    for i in range(retries):
        try:
            r = _api_session.post(govlib.SEARCH_URL, json=body, timeout=30)
            j = r.json()
            res = j.get("result", {})
            data = res.get("data")
            items = data if isinstance(data, list) else (data or {}).get("list", [])
            for it in items:
                u = it.get(govlib.F_URL) or ""
                if isinstance(u, list):
                    u = u[0] if u else ""
                if u == item["url"]:
                    c = it.get("f_202321758948")
                    if isinstance(c, list):
                        c = c[0] if c else ""
                    return str(c or "")
            return ""
        except Exception:
            with _api_lock:
                _api_session = govlib.make_session()
            time.sleep(1 + i)
    return ""


def worker(queue_list, progress, stats, stop_event):
    s = make_session()
    while not stop_event.is_set():
        try:
            item = queue_list.pop(0)
        except IndexError:
            break
        url = item["url"]
        out_dir = os.path.join(DL_ROOT, item["cat"])
        os.makedirs(out_dir, exist_ok=True)
        base = sanitize(f"{item['cat']}_{item['title']}_{item.get('date','')}")
        docx_path = os.path.join(out_dir, base + ".docx")
        html_path = os.path.join(out_dir, base + ".html")

        if os.path.exists(docx_path) and os.path.getsize(docx_path) > 2000:
            progress.mark(url, ok=True, file=os.path.basename(docx_path), skipped=True)
            with stats["lock"]:
                stats["skip"] += 1
            continue

        ok, err = False, ""
        for attempt in range(4):
            try:
                r = s.get(url, timeout=40)
                if r.status_code != 200:
                    raise RuntimeError(f"http {r.status_code}")
                enc = r.apparent_encoding or "utf-8"
                html_text = r.content.decode(enc, errors="replace")
                body = extract_body(html_text)
                if not body or len(body) < 150:
                    body = extract_body_generic(html_text)
                if not body or len(body) < 150:
                    raise RuntimeError(f"正文提取失败 len={len(body)}")
                with open(html_path, "w", encoding="utf-8") as f:
                    f.write(html_text)
                blocks = body_to_blocks(body)
                if len(blocks) < 2:
                    raise RuntimeError(f"块解析过少 {len(blocks)}")
                make_docx(item["title"], blocks, item, docx_path)
                if os.path.getsize(docx_path) < 2000:
                    raise RuntimeError("docx 过小")
                progress.mark(url, ok=True, file=os.path.basename(docx_path),
                              bytes=os.path.getsize(docx_path))
                ok = True
                with stats["lock"]:
                    stats["ok"] += 1
                break
            except Exception as e:
                err = f"{type(e).__name__}: {e}"
                time.sleep(1.5 * (attempt + 1) + random.random())
        if not ok:
            # API 全文回退：从规章库搜索接口按标题取正文字段
            try:
                content = fetch_content_via_api(item)
                if content and len(content) > 300:
                    text = re.sub(r"<[^>]+>", "", content)
                    text = htmllib.unescape(text)
                    paras = [p.strip() for p in re.split(r"\n|　　", text) if p.strip()]
                    make_docx(item["title"], [("p", p) for p in paras], item, docx_path)
                    if os.path.getsize(docx_path) > 2000:
                        progress.mark(url, ok=True,
                                      file=os.path.basename(docx_path),
                                      bytes=os.path.getsize(docx_path),
                                      source="api_content")
                        ok = True
                        with stats["lock"]:
                            stats["ok"] += 1
                        err = ""
            except Exception as e2:
                err += f" | api回退失败: {e2}"
        if not ok:
            progress.mark(url, ok=False, error=err)
            with stats["lock"]:
                stats["fail"] += 1
        with stats["lock"]:
            stats["done"] += 1
            d = stats["done"]
        if d % 100 == 0:
            log(f"[{stats['cat']}] 进度 {d}/{stats['total']} 成功{stats['ok']} "
                f"跳过{stats['skip']} 失败{stats['fail']}")
        time.sleep(0.2 + random.random() * 0.3)


def run(cat, workers=5):
    list_path = os.path.join(LISTS, f"gov_{cat}.jsonl")
    items = []
    with open(list_path, encoding="utf-8") as f:
        for line in f:
            if line.strip():
                items.append(json.loads(line))
    progress = Progress(os.path.join(LISTS, f"gov_dl_progress_{cat}.jsonl"))
    queue_list = [it for it in items if not progress.done_ok(it["url"])]
    log(f"[{cat}] 总 {len(items)}，已完成 {len(items)-len(queue_list)}，待下 {len(queue_list)}")

    stats = {"lock": threading.Lock(), "ok": 0, "fail": 0, "skip": 0, "done": 0,
             "total": len(items), "cat": cat}
    stop_event = threading.Event()
    threads = []
    for i in range(workers):
        t = threading.Thread(target=worker, args=(queue_list, progress, stats, stop_event), daemon=True)
        t.start()
        time.sleep(0.8)
        threads.append(t)
    for t in threads:
        t.join()
    log(f"[{cat}] 结束: 成功{stats['ok']} 跳过{stats['skip']} 失败{stats['fail']} / 共{stats['total']}")


def main():
    cat = sys.argv[1] if len(sys.argv) > 1 else "all"
    workers = int(sys.argv[2]) if len(sys.argv) > 2 else 5
    cats = ["部门规章", "地方政府规章"] if cat == "all" else [cat]
    for c in cats:
        run(c, workers=workers)


if __name__ == "__main__":
    main()
